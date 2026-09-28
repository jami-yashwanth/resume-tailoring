"""Create two sample resumes with their own styling: a DOCX and a PDF with a subset-embedded font."""
import os
import pymupdf
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "samples")
os.makedirs(OUT, exist_ok=True)

BULLETS_RAZORFIN = [
    ("Worked on backend APIs for payments.", None),
    ("Helped refactor the refunds module.", None),
    ("Built a nightly reconciliation job that matches 2.1 lakh transactions against bank files.", "2.1 lakh"),
    ("Mentored 2 junior engineers.", None),
]
BULLETS_KITEBYTE = [
    ("Wrote REST APIs in Java for a loyalty-points service used by 40 merchant partners.", "Java"),
    ("Wrote unit tests.", None),
    ("Moved nightly batch jobs from cron scripts to scheduled Spring tasks.", None),
]


def make_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name = "Georgia"
    base.font.size = Pt(10.5)

    name = doc.add_paragraph()
    r = name.add_run("Priya Sharma")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph("Backend Engineer · Bengaluru · Notice period 30 days · priya.sharma@example.com")

    def heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        p.paragraph_format.space_before = Pt(10)

    def role(left, right):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
        a = p.add_run(left)
        a.bold = True
        p.add_run("\t" + right).italic = True

    def bullet(text, bold_part):
        p = doc.add_paragraph(style="List Bullet")
        if bold_part and bold_part in text:
            before, after = text.split(bold_part, 1)
            p.add_run(before)
            p.add_run(bold_part).bold = True
            p.add_run(after)
        else:
            p.add_run(text)

    heading("Summary")
    doc.add_paragraph("Software engineer interested in backend development.")
    heading("Experience")
    role("Razorfin · Software Engineer, Backend", "Aug 2024 – present")
    for t, b in BULLETS_RAZORFIN:
        bullet(t, b)
    role("Kitebyte Labs · Associate Software Engineer", "Jul 2023 – Jul 2024")
    for t, b in BULLETS_KITEBYTE:
        bullet(t, b)
    heading("Skills")
    doc.add_paragraph("Python, Git, Docker, Java, Spring Boot, AWS, PostgreSQL, Redis")
    heading("Education")
    doc.add_paragraph("B.Tech, Computer Science · CGPA 8.4 · 2023")
    doc.save(path)


def find_font():
    for p in ["/System/Library/Fonts/Supplemental/Georgia.ttf", "/Library/Fonts/Georgia.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"]:
        if os.path.exists(p):
            return p
    return None


def make_pdf(path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)  # A4
    fontfile = find_font()
    page.insert_font(fontname="body", fontfile=fontfile)
    page.insert_font(fontname="bold", fontfile=fontfile.replace("Georgia.ttf", "Georgia Bold.ttf") if fontfile and os.path.exists(fontfile.replace("Georgia.ttf", "Georgia Bold.ttf")) else fontfile)
    ink = (0.106, 0.133, 0.188)
    blue = (0.12, 0.23, 0.37)
    y = 60
    page.insert_text((56, y), "Priya Sharma", fontname="bold", fontsize=20, color=blue); y += 20
    page.insert_text((56, y), "Backend Engineer · Bengaluru · Notice period 30 days · priya.sharma@example.com", fontname="body", fontsize=9.5, color=ink); y += 26

    def heading(t):
        nonlocal y
        page.insert_text((56, y), t.upper(), fontname="bold", fontsize=10.5, color=blue); y += 16

    def line(t, x=56, bold=False):
        nonlocal y
        rect = pymupdf.Rect(x, y - 10, 539, y + 40)
        rc = page.insert_textbox(rect, t, fontname="bold" if bold else "body", fontsize=10, color=ink)
        used = 50 - rc  # height consumed
        y += max(14, int(used) + 2)

    heading("Summary")
    line("Software engineer interested in backend development.")
    y += 6
    heading("Experience")
    line("Razorfin · Software Engineer, Backend        Aug 2024 – present", bold=True)
    for t, _ in BULLETS_RAZORFIN:
        page.insert_text((62, y), "•", fontname="body", fontsize=10, color=ink)
        line(t, x=74)
    line("Kitebyte Labs · Associate Software Engineer        Jul 2023 – Jul 2024", bold=True)
    for t, _ in BULLETS_KITEBYTE:
        page.insert_text((62, y), "•", fontname="body", fontsize=10, color=ink)
        line(t, x=74)
    y += 6
    heading("Skills")
    line("Python, Git, Docker, Java, Spring Boot, AWS, PostgreSQL, Redis")
    y += 6
    heading("Education")
    line("B.Tech, Computer Science · CGPA 8.4 · 2023")
    doc.subset_fonts()  # like real exported PDFs: only the glyphs used are embedded
    doc.save(path, garbage=4, deflate=True)


if __name__ == "__main__":
    make_docx(os.path.join(OUT, "priya_resume.docx"))
    make_pdf(os.path.join(OUT, "priya_resume.pdf"))
    print("samples written to", OUT)
