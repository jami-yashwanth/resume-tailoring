"""Documents that are not shaped like our own sample.

Every bug in here came from a real resume, and every one of them made the
preview wrong while the downloaded file stayed right — the worst failure this
product has, because the screen is what the user checks before they trust it.

The fixture is built in code rather than committed, so the suite never carries
anyone's phone number or email address.
"""
import docx
import pytest
from docx.oxml.ns import qn

from app.docx_ops import has_links, has_numbering, para_text, parse_docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def add_inline_bullet(doc, text):
    """A bullet the way Word actually makes one: numbering applied straight to
    the paragraph, which keeps the style name 'Normal'."""
    p = doc.add_paragraph(text)
    pPr = p._p.get_or_add_pPr()
    numPr = pPr.makeelement(qn("w:numPr"), {})
    ilvl = numPr.makeelement(qn("w:ilvl"), {qn("w:val"): "0"})
    numId = numPr.makeelement(qn("w:numId"), {qn("w:val"): "1"})
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    return p


def add_link(p, text, url):
    """A real w:hyperlink, whose runs are not children of the paragraph."""
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = p._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = p._p.makeelement(qn("w:r"), {})
    t = run.makeelement(qn("w:t"), {})
    t.text = text
    run.append(t)
    link.append(run)
    p._p.append(link)


@pytest.fixture
def real_world(tmp_path):
    doc = docx.Document()

    doc.add_paragraph("Asha Menon")

    contact = doc.add_paragraph("+91 90000 00000 | ")
    add_link(contact, "asha@example.com", "mailto:asha@example.com")
    contact.add_run(" | ")
    add_link(contact, "linkedin.com/in/asha", "https://linkedin.com/in/asha")

    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Northwind Jun 2023 – present")
    add_inline_bullet(doc, "Migrated the Python service to a GPU server.")
    add_inline_bullet(doc, "Upgraded the API from V1 to V2 with stricter validation.")

    path = tmp_path / "real.docx"
    doc.save(str(path))
    return path


def find(blocks, needle):
    return next(b for b in blocks if b.text.startswith(needle))


class TestInlineNumbering:
    def test_a_paragraph_numbered_inline_is_a_bullet(self, real_world):
        """The reported bug: a resume whose bullets are styled 'Normal' with
        inline numbering rendered as flat prose, with no bullets at all."""
        blocks, _, _ = parse_docx(str(real_world))
        migrated = find(blocks, "Migrated the Python service")
        assert migrated.style == "Normal", "precondition: the style says nothing"
        assert migrated.kind.value == "bullet"

    def test_an_ordinary_paragraph_is_still_not_a_bullet(self, real_world):
        blocks, _, _ = parse_docx(str(real_world))
        assert find(blocks, "Northwind").kind.value != "bullet"

    def test_has_numbering_reads_the_paragraph_not_the_style_name(self, real_world):
        document = docx.Document(str(real_world))
        numbered = [p for p in document.paragraphs if has_numbering(p)]
        assert len(numbered) == 2
        assert all(p.style.name == "Normal" for p in numbered)


class TestHyperlinks:
    def test_link_text_is_part_of_the_line(self, real_world):
        """Without this the contact line parses as "+91 90000 00000 |  | ",
        losing the email and the profile entirely."""
        blocks, _, _ = parse_docx(str(real_world))
        contact = find(blocks, "+91")
        assert "asha@example.com" in contact.text
        assert "linkedin.com/in/asha" in contact.text

    def test_link_text_is_carried_into_the_runs_the_preview_draws(self, real_world):
        blocks, _, _ = parse_docx(str(real_world))
        contact = find(blocks, "+91")
        assert "asha@example.com" in "".join(r.text for r in contact.runs)

    def test_a_line_with_a_link_is_flagged(self, real_world):
        blocks, _, _ = parse_docx(str(real_world))
        assert find(blocks, "+91").has_link is True
        assert find(blocks, "Migrated the Python").has_link is False

    def test_no_segment_of_the_contact_line_comes_back_empty(self, real_world):
        """The shape of the reported bug: the separators survived and
        everything between them vanished."""
        document = docx.Document(str(real_world))
        contact = next(p for p in document.paragraphs if has_links(p))
        segments = [s.strip() for s in para_text(contact).split("|")]
        assert len(segments) == 3
        assert all(segments), f"empty segment in {segments}"


class TestDocumentDefaults:
    def test_body_size_falls_back_to_document_defaults(self, real_world):
        """python-docx's default template leaves Normal without an explicit
        size; the size lives in docDefaults. Reading only Normal gave every
        line the same guessed size and flattened the preview's hierarchy."""
        blocks, _, _ = parse_docx(str(real_world))
        assert find(blocks, "Migrated the Python").size > 0

    def test_a_font_is_reported(self, real_world):
        _, fonts, _ = parse_docx(str(real_world))
        assert fonts, "the preview needs to know what the document is set in"


class TestParagraphPresentation:
    """Alignment and rules are how a resume's header and sections read. A
    preview that ignores them puts a centred name flush left and drops the
    line under every heading — visibly not the user's document."""

    @pytest.fixture
    def styled(self, tmp_path):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = docx.Document()
        name = doc.add_paragraph("Asha Menon")
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name.runs[0].font.small_caps = True

        heading = doc.add_paragraph("Experience")
        pPr = heading._p.get_or_add_pPr()
        borders = pPr.makeelement(qn("w:pBdr"), {})
        bottom = borders.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6"})
        borders.append(bottom)
        pPr.append(borders)

        doc.add_paragraph("Plain body text.")
        path = tmp_path / "styled.docx"
        doc.save(str(path))
        return path

    def test_carries_a_centred_header(self, styled):
        blocks, _, _ = parse_docx(str(styled))
        assert find(blocks, "Asha Menon").align == "center"
        assert find(blocks, "Plain body").align == "left"

    def test_carries_small_caps(self, styled):
        blocks, _, _ = parse_docx(str(styled))
        assert find(blocks, "Asha Menon").runs[0].small_caps is True

    def test_carries_the_rule_under_a_heading(self, styled):
        blocks, _, _ = parse_docx(str(styled))
        assert find(blocks, "Experience").rule_below is True
        assert find(blocks, "Plain body").rule_below is False
